"""Procesamiento por lotes de CobDosel (sin dependencias de Qt ni de QGIS).

Diseño de ejemplo (cualquier diseño sirve: el formulario indica parcela y punto de cada foto):
- parcelas circulares de 500 m² (radio 12.62 m);
- 5 fotos verticales por parcela: Centro y N-O (315°), N-E (45°), S-E (135°), S-O (225°) a 8.9 m
  del centro;
- valor de la parcela = media de las fotos válidas disponibles;
- valor del sitio = media de las medias de parcela (Índice CaCo).

Entradas: carpeta de fotos (nombres de la cámara), formulario de campo digital (CSV o Excel con las
columnas de la plantilla), calibración guardada, máscara de horizonte (fija, 63° por defecto, el estándar
de CobDosel; Tichý 2016 usó 60°; u opcionalmente desde una altura del dosel del sitio) y el número de parcelas planificadas.
El protocolo no mide alturas: los árboles se asignan a grupos de dosel (CANOPY_GROUPS). El nivel de corte de cada foto es el recomendado
(cruce de dos log-normales, método de la app GLAMA; iterado hasta estabilizarse desde 128); si no hay
recomendación se usa Otsu.

Autor: Jorge Fallas (jfallas56@gmail.com)
"""

import csv
import datetime
import math
import os
import re
import unicodedata

import numpy as np

from .exif_gps import load_oriented_rgb_image, extract_exif
from .focal_check import calibration_focal_mm
from .camera_check import camera_fingerprint, camera_label, camera_group_key, compare_cameras
from .lens_calibration import scale_calibration_to_image, zenith_to_pixel_radius
from .canopy_analysis import (
    extract_channel, threshold_image, compute_otsu_threshold, compute_radius_grid,
    compute_zenith_grid, analysis_circle_visibility,
)
from .glama_engine import (  # noqa: F401  (glama_cut_iterated se reexporta por compatibilidad)
    glama_ring_counts, glama_ring_counts_with_gray, glama_indices, glama_recommended_cut,
    glama_cut_iterated, GLAMA_CUT_START, glama_cut_diagnosis,
)
import json

PLOT_RADIUS_M = 12.62            # parcela circular de 500 m²
DEFAULT_FOOTPRINT_M = 6.31       # radio equivalente del cuadrante (125 m²) cubierto por cada foto
DEFAULT_CAMERA_HEIGHT_M = 1.3
DEFAULT_TOLERANCE_S = 60.0
DEFAULT_MASK_DEG = 63.0          # v0.14.29: estándar de CobDosel (cenit ≤ 27°): círculo completo con teléfonos de 23–25 mm; Tichý (2016) usó 60°

# Grupos de dosel del protocolo (posición del árbol; no se mide la altura). (código, nombre, min m, max m)
CANOPY_GROUPS = (
    (1, "Rastrero: herbáceas y helechos", 0.0, 1.0),
    (2, "Sotobosque: regeneración y arbustos", 1.0, 5.0),
    (3, "Copa inferior / Subdosel", 5.0, 15.0),
    (4, "Copa superior / Dosel principal", 15.0, 30.0),
    (5, "Copas emergentes (no compactas)", 30.0, None),
)


def canopy_group_height(code):
    """Altura de referencia de un grupo de dosel (punto medio; límite inferior en el grupo 5)."""
    for c, _name, lo, hi in CANOPY_GROUPS:
        if c == int(code):
            return lo if hi is None else (lo + hi) / 2.0
    raise ValueError("Grupo de dosel desconocido: {}".format(code))


MAX_ANALYSIS_DIM = 4096
IMAGE_EXT = (".jpg", ".jpeg", ".png", ".tif", ".tiff", ".webp", ".heic", ".heif")

# Puntos del diseño de ejemplo (orden del formulario) y sus sinónimos aceptados
POINTS = ("Centro", "N-O", "N-E", "S-E", "S-O")
_POINT_ALIASES = {
    "c": "Centro", "centro": "Centro", "center": "Centro", "0": "Centro",
    "no": "N-O", "nw": "N-O", "315": "N-O", "q4": "N-O",
    "ne": "N-E", "45": "N-E", "q1": "N-E",
    "se": "S-E", "135": "S-E", "q2": "S-E",
    "so": "S-O", "sw": "S-O", "225": "S-O", "q3": "S-O",
}

# Columnas de la plantilla y sinónimos (comparación sin tildes ni mayúsculas)
FORM_COLUMNS = ("parcela", "reemplaza_a", "punto", "distancia_m",
                "id_foto", "hora_captura", "calidad", "observaciones")
_COLUMN_ALIASES = {
    "parcela": ("parcela", "plot", "id_parcela", "n_parcela"),
    "reemplaza_a": ("reemplaza_a", "reemplaza", "reemplazo_de", "reemplaza_a_parcela"),
    "punto": ("punto", "posicion", "point"),
    "distancia_m": ("distancia_m", "distancia", "distancia_diagonal_m", "distancia_diagonal"),
    "id_foto": ("id_foto", "id_de_fotografia", "id_fotografia", "foto", "archivo"),
    "hora_captura": ("hora_captura", "hora_de_captura", "fecha_hora", "hora"),
    "calidad": ("calidad", "estado"),
    "observaciones": ("observaciones", "obs", "comentarios"),
}
# Calidades que excluyen la foto del cálculo. v0.14.38: se compara el valor completo (o cada parte
# separada por «,», «;» o «/»), no fragmentos: «sin sol» o «buena (no repetir)» ya no excluyen la foto.
_BAD_QUALITY = ("sol", "sol directo", "movida", "foto movida", "desenfocada", "repetir", "mala",
                "descartar", "excluir", "excluida", "no valida", "invalida")


def quality_excludes(value):
    """True si la calidad anotada excluye la foto (valor completo o una de sus partes)."""
    parts = re.split(r"[,;/]", str(value or ""))
    return any(_norm(p).replace("_", " ") in _BAD_QUALITY for p in parts if p.strip())


def _norm(text):
    """minúsculas, sin tildes, espacios y signos → «_»."""
    t = unicodedata.normalize("NFKD", str(text or "")).encode("ascii", "ignore").decode("ascii")
    return re.sub(r"[^a-z0-9]+", "_", t.lower()).strip("_")


def normalize_point(value):
    """«N-O (315°)», «NO», «nw», «315» → «N-O»; «Centro», «C» → «Centro». None si no se reconoce."""
    key = _norm(value).replace("_", "")
    if key in _POINT_ALIASES:
        return _POINT_ALIASES[key]
    for k in ("centro", "no", "ne", "se", "so", "nw", "sw"):
        if key.startswith(k):
            return _POINT_ALIASES[k]
    digits = re.findall(r"\d+", key)
    if digits and digits[0] in _POINT_ALIASES:
        return _POINT_ALIASES[digits[0]]
    return None


# ----------------------------------------------------------------------
# Fechas y horas
# ----------------------------------------------------------------------
_DT_PATTERNS = (
    "%d/%m/%Y %H:%M:%S", "%d/%m/%Y %H:%M", "%d-%m-%Y %H:%M:%S", "%d-%m-%Y %H:%M",
    "%Y-%m-%d %H:%M:%S", "%Y-%m-%d %H:%M", "%Y:%m:%d %H:%M:%S", "%d/%m/%y %H:%M:%S",
    "%d/%m/%y %H:%M",
)


def parse_datetime(value, default_date=None):
    """
    Fecha y hora del formulario, día primero (DD/MM/AAAA). Acepta «17/10/2026 / 19:49:43»,
    «17/10/2026 19:49», «2026-10-17T19:49:43», objetos datetime y, con default_date, solo la hora.
    Devuelve datetime o None.
    """
    if value is None or value == "":
        return None
    if isinstance(value, datetime.datetime):
        return value
    if isinstance(value, datetime.time) and default_date is not None:
        return datetime.datetime.combine(default_date, value)
    s = str(value).strip().replace("T", " ")
    s = re.sub(r"\s*/\s*(?=\d{1,2}:)", " ", s)      # «17/10/2026 / 19:49:43» → «17/10/2026 19:49:43»
    s = re.sub(r"\s+", " ", s)
    for fmt in _DT_PATTERNS:
        try:
            return datetime.datetime.strptime(s, fmt)
        except ValueError:
            continue
    if default_date is not None:
        for fmt in ("%H:%M:%S", "%H:%M"):
            try:
                return datetime.datetime.combine(default_date, datetime.datetime.strptime(s, fmt).time())
            except ValueError:
                continue
    return None


# ----------------------------------------------------------------------
# Formulario de campo
# ----------------------------------------------------------------------
def _map_header(header):
    mapping = {}
    normed = [_norm(h) for h in header]
    for key, aliases in _COLUMN_ALIASES.items():
        for i, h in enumerate(normed):
            if h in aliases and i not in mapping.values():
                mapping[key] = i
                break
    return mapping


def _read_table(path):
    ext = os.path.splitext(path)[1].lower()
    if ext in (".csv", ".txt"):
        with open(path, "r", encoding="utf-8-sig", newline="") as f:
            sample = f.read(4096)
            f.seek(0)
            try:
                dialect = csv.Sniffer().sniff(sample, delimiters=",;\t")
            except csv.Error:
                dialect = csv.excel
            return [row for row in csv.reader(f, dialect)]
    if ext in (".xlsx", ".xlsm"):
        try:
            import openpyxl
            wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
            ws = wb.worksheets[0]
            rows = [[c for c in r] for r in ws.iter_rows(values_only=True)]
            wb.close()
            return rows
        except ImportError:
            pass
        try:
            from osgeo import ogr      # disponible en QGIS
        except ImportError:
            raise ValueError("Para leer Excel se necesita openpyxl o GDAL; guarde el formulario como CSV.")
        ds = ogr.Open(path)
        if ds is None:
            raise ValueError("No se pudo abrir el formulario: " + path)
        lyr = ds.GetLayer(0)
        defn = lyr.GetLayerDefn()
        header = [defn.GetFieldDefn(i).GetName() for i in range(defn.GetFieldCount())]
        rows = [header]
        for feat in lyr:
            rows.append([feat.GetField(i) for i in range(defn.GetFieldCount())])
        return rows
    raise ValueError("Formato de formulario no soportado (use .csv o .xlsx): " + path)


def read_field_form(path):
    """
    Lee el formulario de fotos. Devuelve (filas, avisos). Cada fila: dict con las claves de
    FORM_COLUMNS, más «fila» (n.º de fila en el archivo), «punto_norm» y «hora_dt».
    Las celdas vacías de parcela y reemplaza_a heredan el valor de la fila anterior (el formulario en
    papel solo las anota una vez por parcela). Otras columnas (p. ej. una altura de 0.12.0) se ignoran.
    """
    table = _read_table(path)
    header_idx = None
    for i, row in enumerate(table[:20]):
        mp = _map_header(row)
        if "parcela" in mp and ("id_foto" in mp or "hora_captura" in mp):
            header_idx, mapping = i, mp
            break
    if header_idx is None:
        raise ValueError("No se encontró la fila de encabezados (se esperan al menos «parcela» e "
                         "«id_foto» o «hora_captura»).")
    rows, warnings = [], []
    carry = {"parcela": None, "reemplaza_a": None}
    for n, raw in enumerate(table[header_idx + 1:], start=header_idx + 2):
        if raw is None or all((c is None or str(c).strip() == "") for c in raw):
            continue
        # Fila descriptiva de la plantilla («Parcela», «ID de fotografía», …): se omite
        if len(_map_header(raw)) >= 3 or _norm(raw[mapping["parcela"]]) == "parcela":
            continue
        rec = {}
        for key in FORM_COLUMNS:
            idx = mapping.get(key)
            val = raw[idx] if idx is not None and idx < len(raw) else None
            rec[key] = val.strip() if isinstance(val, str) else val
        if rec["parcela"] not in (None, ""):
            # Nueva parcela: sus datos de cabecera valen hasta la próxima parcela
            carry = {"parcela": rec["parcela"], "reemplaza_a": rec["reemplaza_a"]}
        else:
            rec["parcela"] = carry["parcela"]
            if rec["reemplaza_a"] in (None, ""):
                rec["reemplaza_a"] = carry["reemplaza_a"]
            else:
                carry["reemplaza_a"] = rec["reemplaza_a"]
        if rec["parcela"] in (None, ""):
            warnings.append("Fila {}: sin parcela; se omite.".format(n))
            continue
        rec["parcela"] = str(rec["parcela"]).strip()
        if re.fullmatch(r"\d+\.0", rec["parcela"]):        # 7.0 leído de Excel → «7»
            rec["parcela"] = rec["parcela"][:-2]
        rec["fila"] = n
        rec["punto_norm"] = normalize_point(rec["punto"]) if rec["punto"] not in (None, "") else None
        if rec["punto"] not in (None, "") and rec["punto_norm"] is None:
            warnings.append("Fila {}: punto «{}» no reconocido.".format(n, rec["punto"]))
        rec["hora_dt"] = parse_datetime(rec["hora_captura"])
        if rec["hora_captura"] not in (None, "") and rec["hora_dt"] is None:
            warnings.append("Fila {}: hora «{}» no reconocida (use DD/MM/AAAA HH:MM:SS).".format(
                n, rec["hora_captura"]))
        rec["excluida_por_calidad"] = quality_excludes(rec["calidad"])
        rows.append(rec)
    return rows, warnings


# ----------------------------------------------------------------------
# Fotos y emparejamiento
# ----------------------------------------------------------------------
def photo_datetime(path):
    """Fecha y hora EXIF (DateTimeOriginal o DateTime) de la foto, o None."""
    info = extract_exif(path)
    return parse_datetime(info.get("capture_datetime"))


def _photo_camera(path):
    """Huella de cámara de una foto (EXIF + tamaño del archivo; la proporción no depende del giro)."""
    from PIL import Image
    try:
        with Image.open(path) as im:
            w, h = im.size
    except Exception:
        return None
    return camera_fingerprint(extract_exif(path), w, h)


def scan_photos(folder):
    """Lista de dicts {ruta, nombre, base, hora_exif, camara} de las imágenes de la carpeta (no recursivo)."""
    out = []
    for name in sorted(os.listdir(folder)):
        p = os.path.join(folder, name)
        if os.path.isfile(p) and name.lower().endswith(IMAGE_EXT):
            out.append({"ruta": p, "nombre": name, "base": os.path.splitext(name)[0],
                        "hora_exif": photo_datetime(p), "camara": _photo_camera(p)})
    return out


_POINT_AZIMUTH_LABEL = {"N-O": " (315°)", "N-E": " (45°)", "S-E": " (135°)", "S-O": " (225°)"}


def draft_form_from_photos(folder, out_path, photos_per_plot=5, template_xlsx=None):
    """
    Crea un formulario de fotos prellenado con todas las fotos de la carpeta, ordenadas por hora EXIF
    (o por nombre si no hay hora): id_foto y hora_captura escritos; parcela y punto asignados por orden,
    photos_per_plot fotos por parcela (con 5: Centro, N-O, N-E, S-E, S-O). El usuario debe revisarlo.
    Escribe .xlsx (desde la plantilla, con openpyxl) o .csv. Devuelve (ruta, n.º de fotos, avisos).
    """
    photos = scan_photos(folder)
    if not photos:
        raise ValueError("La carpeta no contiene fotos ({}).".format(", ".join(IMAGE_EXT)))
    far = datetime.datetime(9999, 1, 1)
    photos.sort(key=lambda p: (p["hora_exif"] or far, p["nombre"].lower()))
    n = max(1, int(photos_per_plot))
    warnings = []
    no_time = [p["nombre"] for p in photos if p["hora_exif"] is None]
    if no_time:
        warnings.append("{} foto(s) sin hora EXIF, al final por nombre: {}.".format(
            len(no_time), ", ".join(no_time[:5]) + (" …" if len(no_time) > 5 else "")))
    rows = []
    for i, ph in enumerate(photos):
        k = i % n
        if n == len(POINTS):
            punto = POINTS[k] + _POINT_AZIMUTH_LABEL.get(POINTS[k], "")
            dist = 0.0 if k == 0 else 8.9
        else:
            punto, dist = "Foto {}".format(k + 1), ""
        hora = ph["hora_exif"].strftime("%d/%m/%Y / %H:%M:%S") if ph["hora_exif"] else ""
        rows.append([str(i // n + 1), "", punto, dist, ph["base"], hora, "OK",
                     "asignación automática: verificar"])
    if len(photos) % n:
        warnings.append("{} fotos no es múltiplo de {}: la última parcela queda incompleta.".format(
            len(photos), n))
    # Saltos de tiempo dentro de una parcela (posible asignación errónea)
    for start in range(0, len(photos), n):
        grp = [p["hora_exif"] for p in photos[start:start + n] if p["hora_exif"]]
        if len(grp) > 1 and (max(grp) - min(grp)).total_seconds() > 20 * 60:
            warnings.append("Parcela {}: sus fotos abarcan más de 20 min; revise la asignación.".format(
                start // n + 1))
    if out_path.lower().endswith(".xlsx"):
        import openpyxl
        tpl = template_xlsx or os.path.join(os.path.dirname(__file__), "plantillas",
                                            "Formulario_fotos_CobDosel.xlsx")
        wb = openpyxl.load_workbook(tpl)
        ws = wb.worksheets[0]
        if ws.max_row >= 3:
            ws.delete_rows(3, ws.max_row)
        for r in rows:
            ws.append(r)
        wb.save(out_path)
    else:
        with open(out_path, "w", newline="", encoding="utf-8-sig") as f:
            w = csv.writer(f)
            w.writerow(FORM_COLUMNS)
            w.writerows(rows)
    return out_path, len(photos), warnings


def camera_groups(photos):
    """[(etiqueta, n.º de fotos, huella)] por lente/encuadre, de más a menos fotos."""
    groups = {}
    for ph in photos:
        cam = ph.get("camara")
        key = camera_group_key(cam)
        g = groups.setdefault(key, {"n": 0, "cam": cam, "sizes": set()})
        g["n"] += 1
        if cam and cam.get("size"):
            g["sizes"].add("{}×{}".format(*cam["size"]))
    out = []
    for g in sorted(groups.values(), key=lambda v: -v["n"]):
        base = camera_label(dict(g["cam"] or {}, size=None))
        out.append(("{} [{}]".format(base, ", ".join(sorted(g["sizes"])) or "—"), g["n"], g["cam"]))
    return out


def match_photos(rows, photos, tolerance_s=DEFAULT_TOLERANCE_S, clock_offset_s=0.0):
    """
    Empareja cada fila del formulario con una foto: primero por ID (nombre de archivo, sin
    extensión ni mayúsculas); si no hay ID o no aparece, por la hora EXIF más cercana dentro de la
    tolerancia. clock_offset_s se suma a la hora EXIF (corrige el reloj de la cámara).
    Devuelve (pares, avisos); cada par = (fila, foto o None, método, diferencia_s o None).
    """
    by_id = {}
    for ph in photos:
        by_id.setdefault(ph["base"].lower(), ph)
        by_id.setdefault(ph["nombre"].lower(), ph)
    used = set()
    pairs, warnings = [], []
    pending = []
    for r in rows:
        pid = str(r.get("id_foto") or "").strip()
        ph = None
        if pid:
            key = pid.lower()
            ph = by_id.get(key) or by_id.get(os.path.splitext(key)[0])
            if ph is None:
                cands = [p for p in photos if key in p["base"].lower()]
                ph = cands[0] if len(cands) == 1 else None
        if ph is not None and ph["ruta"] not in used:
            diff = None
            if r.get("hora_dt") and ph.get("hora_exif"):
                diff = (ph["hora_exif"] + datetime.timedelta(seconds=clock_offset_s)
                        - r["hora_dt"]).total_seconds()
                if abs(diff) > tolerance_s:
                    warnings.append(
                        "Parcela {} {}: la foto {} coincide por ID, pero su hora difiere {:+.0f} s del "
                        "formulario (revise).".format(r["parcela"], r.get("punto_norm") or "", ph["nombre"], diff))
            used.add(ph["ruta"])
            pairs.append((r, ph, "id", diff))
        else:
            if pid and ph is None:
                warnings.append("Parcela {} {}: ID «{}» no encontrado en la carpeta; se intenta por hora."
                                .format(r["parcela"], r.get("punto_norm") or "", pid))
            pending.append(r)
    for r in pending:
        best, best_d = None, None
        if r.get("hora_dt"):
            for ph in photos:
                if ph["ruta"] in used or ph.get("hora_exif") is None:
                    continue
                d = (ph["hora_exif"] + datetime.timedelta(seconds=clock_offset_s) - r["hora_dt"]).total_seconds()
                if abs(d) <= tolerance_s and (best_d is None or abs(d) < abs(best_d)):
                    best, best_d = ph, d
        if best is not None:
            used.add(best["ruta"])
            pairs.append((r, best, "hora", best_d))
        else:
            warnings.append("Parcela {} {} (fila {}): sin foto correspondiente.".format(
                r["parcela"], r.get("punto_norm") or "", r["fila"]))
            pairs.append((r, None, "sin foto", None))
    for ph in photos:
        if ph["ruta"] not in used:
            warnings.append("Foto {} sin fila en el formulario (no se usa).".format(ph["nombre"]))
    pairs.sort(key=lambda t: t[0]["fila"])
    return pairs, warnings


# ----------------------------------------------------------------------
# Máscara desde la altura del dosel
# ----------------------------------------------------------------------
def zenith_from_height(canopy_height_m, camera_height_m=DEFAULT_CAMERA_HEIGHT_M,
                       footprint_radius_m=DEFAULT_FOOTPRINT_M):
    """Ángulo cenital límite θ = atan(radio de la huella / (altura del dosel − altura de la cámara))."""
    h = float(canopy_height_m) - float(camera_height_m)
    if h <= 0:
        raise ValueError("La altura del dosel debe ser mayor que la altura de la cámara.")
    return math.degrees(math.atan(float(footprint_radius_m) / h))


# ----------------------------------------------------------------------
# Análisis de una foto
# ----------------------------------------------------------------------
def analyze_photo(path, calibration, zenith_max_deg, channel_mode="blue_red", sky_is_bright=True,
                  cut_rule="glama", fixed_cut=None):
    """
    Analiza una foto con la calibración dada y la máscara (cenit máximo, grados). Devuelve dict con
    caco, modif_caco, apertura, cierre, fraccion_claros, corte, metodo_corte, visible_pct, tamaño y avisos.
    """
    # v0.14.38: sin parámetros ignorados en silencio
    if cut_rule == "fixed" and fixed_cut is None:
        raise ValueError("cut_rule='fixed' requiere fixed_cut.")
    if cut_rule != "fixed" and fixed_cut is not None:
        raise ValueError("fixed_cut solo se usa con cut_rule='fixed' (se recibió cut_rule={!r}).".format(cut_rule))
    if cut_rule not in ("glama", "otsu", "fixed"):
        raise ValueError("cut_rule desconocido: {!r}".format(cut_rule))
    img, _scale, orig = load_oriented_rgb_image(path, max_dim=MAX_ANALYSIS_DIM)
    arr = np.asarray(img)
    h, w = arr.shape[:2]
    cal = scale_calibration_to_image(calibration, w, h)
    R = cal.get("virtual_radius_px")
    if not R:
        raise ValueError("La calibración no define el radio del hemisferio.")
    # v0.14.23: centro de la calibración (como en el análisis de una foto); antes siempre el centro de
    # la imagen, distinto del diálogo con un ojo de pez cuyo círculo detectado no está centrado.
    cc = cal.get("center_xy")
    c = (float(cc[0]), float(cc[1])) if cc is not None else (w / 2.0, h / 2.0)
    proj = cal.get("projection", "polar")
    gray = extract_channel(arr, channel_mode)
    warns = []
    if cal.pop("center_note", None):
        warns.append("foto girada respecto de la calibración con centro desplazado: se usó el centro de la imagen")
    meas = calibration.get("measurement") or {}
    ref = meas.get("image_size")
    if ref:
        a_ref = max(ref) / float(min(ref))
        a_now = max(w, h) / float(min(w, h))
        if abs(a_ref - a_now) > 0.01 * a_now:
            warns.append("proporción de la foto distinta de la calibración")
    method = cut_rule
    # Corte automático (recomendado log-normal) y comprobación del método: siempre se registra
    diag = {"cut": None, "ok": False, "reasons": ["cielo oscuro: sin corte automático"], "starts": {}}
    if sky_is_bright:
        hist = glama_ring_counts_with_gray(gray, GLAMA_CUT_START, c, R, zenith_max_deg,
                                           calibration=cal, projection=proj)["histogram"]
        diag = glama_cut_diagnosis(hist)
    auto_cut = diag["cut"]
    revisar = False
    if cut_rule == "fixed":
        t = int(fixed_cut)
        method = "fijo"          # el lote lo cambia a «revisado por el operador» si viene de la revisión
    elif cut_rule == "otsu":
        rp = compute_radius_grid((h, w), c)
        area = compute_zenith_grid(rp, cal) <= zenith_max_deg
        t = compute_otsu_threshold(gray, area)
    else:
        revisar = not diag["ok"]
        if auto_cut is None:
            rp = compute_radius_grid((h, w), c)
            area = compute_zenith_grid(rp, cal) <= zenith_max_deg
            t = compute_otsu_threshold(gray, area)
            method = "otsu provisional (sin recomendación log-normal)"
        else:
            t = auto_cut
            method = "recomendado (log-normal)"
        if revisar:
            warns.append("revisar el corte: " + "; ".join(diag["reasons"]))
    sky = threshold_image(gray, t, sky_is_bright=sky_is_bright)
    counts = glama_ring_counts(sky, c, R, zenith_max_deg, calibration=cal, projection=proj)
    idx = glama_indices(counts, calibration=cal, radius_px=R)
    try:
        vis = analysis_circle_visibility((h, w), c, float(zenith_to_pixel_radius(zenith_max_deg, cal)))
    except (ValueError, KeyError, ZeroDivisionError):
        vis = None
    if vis is not None and vis < 99.5:
        warns.append("círculo de análisis visible {:.1f} %".format(vis))
    return {
        "caco": idx["canopy_cover_index_pct"],
        "modif_caco": idx["spherical_cover_pct"],
        "apertura": idx["canopy_openness_pct"],
        "cierre": idx["canopy_closure_pct"],
        "fraccion_claros": idx["gap_fraction_pct"],
        "corte": int(t),
        "metodo_corte": method,
        "corte_auto": auto_cut,
        "revisar": revisar,
        "motivo_revision": "; ".join(diag["reasons"]) if not diag["ok"] else "",
        "inicios": "; ".join("{}→{}".format(k, v) for k, v in diag.get("starts", {}).items()),
        "visible_pct": vis,
        "ancho": w, "alto": h, "original": orig,
        "avisos": warns,
    }


# ----------------------------------------------------------------------
# Lote completo
# ----------------------------------------------------------------------
def run_batch(photo_folder, form_path, calibration, planned_plots, mask_mode="fixed",
              canopy_height_m=None, camera_height_m=DEFAULT_CAMERA_HEIGHT_M,
              footprint_radius_m=DEFAULT_FOOTPRINT_M, fixed_mask_deg=DEFAULT_MASK_DEG, channel_mode="blue_red",
              sky_is_bright=True, cut_rule="glama", fixed_cut=None,
              tolerance_s=DEFAULT_TOLERANCE_S, clock_offset_s=0.0, progress=None, log=None,
              canopy_group=None, decisions=None):
    """
    Ejecuta el lote. mask_mode: «fixed» (fixed_mask_deg sobre el horizonte; 63° por defecto) o
    «height» (máscara desde canopy_height_m, una altura de referencia del sitio, p. ej. la de un grupo
    de dosel). decisions: decisiones de la revisión visual de cortes {archivo: {decision, corte, ...}}
    (ver load_decisions). Devuelve dict con fotos, parcelas, sitio, avisos y parámetros.
    """
    decisions = decisions or {}
    if mask_mode == "height":
        if canopy_height_m is None:
            raise ValueError("Máscara desde la altura: indique la altura del dosel del sitio.")
        zen_height = zenith_from_height(canopy_height_m, camera_height_m, footprint_radius_m)
    elif mask_mode == "fixed":
        if fixed_mask_deg is None or not 0.0 <= float(fixed_mask_deg) < 90.0:
            raise ValueError("Máscara fija no válida (0–90° sobre el horizonte).")
    else:
        raise ValueError("mask_mode desconocido: {}".format(mask_mode))

    def _log(msg, level="INFO"):
        if callable(log):
            log(msg, level)

    rows, warnings = read_field_form(form_path)
    photos = scan_photos(photo_folder)
    pairs, w2 = match_photos(rows, photos, tolerance_s, clock_offset_s)
    warnings += w2
    # Cámaras/lentes: una calibración vale para una lente y un encuadre (no para otra lente o zoom)
    cal_cam = calibration.get("camera") or (calibration.get("measurement") or {}).get("camera")
    if cal_cam and not (cal_cam.get("model") or cal_cam.get("focal_length_mm") or cal_cam.get("focal_length_35mm")):
        cal_cam = None          # huella sin EXIF útil (foto de calibración sin modelo ni focal)
    used = [ph for _r, ph, _h, _d in pairs if ph is not None]
    groups = camera_groups(used)
    if groups:
        warnings.append("Fotos por cámara/lente: " + "; ".join("{}: {}".format(lbl, n) for lbl, n, _c in groups))
    if not cal_cam and len(groups) > 1:
        raise ValueError(
            "Las fotos son de {} cámaras/lentes distintas y la calibración no registra la cámara:\n• {}\n"
            "Separe las fotos por cámara y procese cada grupo con su calibración, o vuelva a guardar la "
            "calibración en CobDosel ≥ 0.12.4 con una foto de esa cámara.".format(
                len(groups), "\n• ".join("{}: {} fotos".format(lbl, n) for lbl, n, _c in groups)))
    if not cal_cam and len(groups) == 1:
        warnings.append("La calibración no registra la cámara. Todas las fotos son de una misma lente ({}): "
                        "verifique que sea la de la calibración y regístrela con «Registrar cámara…».".format(
                            groups[0][0]))
    elif not cal_cam:
        warnings.append("La calibración no registra la cámara: no se pudo verificar la lente de las fotos.")
    results = []
    n = len(pairs)
    for i, (r, ph, how, diff) in enumerate(pairs):
        rec = {"parcela": r["parcela"], "reemplaza_a": r.get("reemplaza_a") or "",
               "punto": r.get("punto_norm") or (r.get("punto") or ""), "fila": r["fila"],
               "id_foto": r.get("id_foto") or "", "hora_formulario": r.get("hora_captura") or "",
               "archivo": ph["nombre"] if ph else "", "emparejado_por": how,
               "dif_hora_s": diff,
               "hora_exif": ph["hora_exif"].isoformat(sep=" ") if ph and ph.get("hora_exif") else "",
               "calidad": r.get("calidad") or "", "valida": False, "motivo": "",
               "camara": camera_label(ph.get("camara")) if ph else "",
               "tamano_px": "{}×{}".format(*ph["camara"]["size"]) if ph and ph.get("camara") else "",
               "mascara_horizonte": None, "cenit_max": None}
        try:
            if ph is None:
                rec["motivo"] = "sin foto"
            elif r.get("excluida_por_calidad"):
                rec["motivo"] = "excluida por calidad («{}»)".format(r.get("calidad"))
            elif cal_cam and compare_cameras(cal_cam, ph.get("camara"))[0] is False:
                rec["motivo"] = "calibración de otra cámara/lente: " + "; ".join(
                    compare_cameras(cal_cam, ph.get("camara"))[1])
            else:
                # v0.14.23: foto sin EXIF comparable → se analiza, pero se deja constancia en la fila
                cam_unverified = bool(cal_cam) and compare_cameras(cal_cam, ph.get("camara"))[0] is None
                zen = 90.0 - float(fixed_mask_deg) if mask_mode == "fixed" else zen_height
                rec["cenit_max"] = zen
                rec["mascara_horizonte"] = 90.0 - zen
                dec = decisions.get(ph["nombre"])
                if dec and dec.get("canal") and dec.get("canal") != channel_mode:
                    warnings.append("{}: la revisión se hizo con otro canal ({}); se ignora.".format(
                        ph["nombre"], dec.get("canal")))
                    dec = None
                rec.update({"decision": "", "operador": "", "nota": ""})
                if dec and dec.get("decision") == "excluido":
                    rec.update({"decision": "excluido", "operador": dec.get("operador", ""),
                                "nota": dec.get("nota", ""), "corte_auto": dec.get("corte_auto")})
                    rec["motivo"] = "excluida por el operador" + (
                        " ({})".format(dec["nota"]) if dec.get("nota") else "")
                else:
                    use_fixed = dec and dec.get("decision") in ("aceptado", "ajustado")
                    res = analyze_photo(ph["ruta"], calibration, zen, channel_mode, sky_is_bright,
                                        "fixed" if use_fixed else cut_rule,
                                        int(dec["corte"]) if use_fixed else (fixed_cut if cut_rule == "fixed" else None))
                    rec.update(res)
                    rec["valida"] = True
                    if use_fixed:
                        rec["decision"] = dec["decision"] if dec["decision"] == "aceptado" or (
                            res.get("corte_auto") == int(dec["corte"])) else "ajustado"
                        if rec["decision"] == "ajustado":
                            rec["decision"] = "ajustado ({} → {})".format(res.get("corte_auto"), int(dec["corte"]))
                        rec["operador"], rec["nota"] = dec.get("operador", ""), dec.get("nota", "")
                        rec["metodo_corte"] = "revisado por el operador"
                        rec["revisar"] = False
                        res["avisos"] = [a for a in res["avisos"] if not a.startswith("revisar el corte")]
                    elif cut_rule == "glama":
                        rec["decision"] = "pendiente de revisión" if res.get("revisar") else "automático"
                    if cam_unverified:
                        res["avisos"].append("cámara no verificada (la foto no tiene modelo ni distancia focal en el EXIF)")
                    rec["motivo"] = "; ".join(res["avisos"])
        except Exception as e:      # una foto defectuosa no detiene el lote
            rec["motivo"] = "error: {}".format(e)
            _log("{}: {}".format(rec.get("archivo") or rec["parcela"], e), "WARN")
        results.append(rec)
        if callable(progress):
            progress(i + 1, n)
    summary = aggregate(results, planned_plots)
    return {"fotos": results, "parcelas": summary["parcelas"], "sitio": summary["sitio"],
            "avisos": sorted(warnings + summary["avisos"], key=lambda w: not w.startswith("Resultado provisional")),
            "parametros": {"carpeta": photo_folder, "formulario": form_path,
                           "calibracion": calibration.get("lens_name") or calibration.get("projection_label"),
                           "camara_calibracion": camera_label(cal_cam) if cal_cam else "no registrada",
                           # v0.14.27: la calibración completa queda registrada con el resultado
                           **calibration_record(calibration),
                           "mascara_modo": "fija" if mask_mode == "fixed" else "altura del dosel",
                           "mascara_horizonte_deg": (float(fixed_mask_deg) if mask_mode == "fixed"
                                                     else round(90.0 - zen_height, 2)),
                           "grupo_dosel": canopy_group if mask_mode == "height" else None,
                           "altura_dosel_sitio_m": canopy_height_m if mask_mode == "height" else None,
                           "altura_camara_m": camera_height_m if mask_mode == "height" else None,
                           "radio_huella_m": footprint_radius_m if mask_mode == "height" else None,
                           "canal": channel_mode, "regla_corte": cut_rule,
                           "tolerancia_s": tolerance_s, "desfase_reloj_s": clock_offset_s,
                           "parcelas_planificadas": planned_plots}}


def calibration_record(calibration):
    """Datos de la calibración para el registro del lote (CSV del sitio e informe HTML)."""
    cal = calibration or {}
    meta = cal.get("calibration_metadata") or {}
    d = cal.get("circle_diameter_cm") or meta.get("circle_diameter_cm")
    h = cal.get("camera_height_cm") or meta.get("camera_height_cm")
    method = cal.get("calibration_method")
    if method == "glama_radius":
        method = "radio de GLAMA ({:g} px)".format(float(cal.get("glama_radius_px") or 0))
    elif not method:
        method = ("círculo D/H" if d else
                  "detección del círculo de 180°" if cal.get("lens_type") == "fisheye" else "—")
    w, hh = cal.get("image_width_px"), cal.get("image_height_px")
    fmm = calibration_focal_mm(dict(cal, circle_diameter_cm=d, camera_height_cm=h))
    return {
        "calibracion_modelo": cal.get("projection_label") or cal.get("projection"),
        "calibracion_metodo": method,
        "circulo_D_cm": round(float(d), 2) if d else None,
        "altura_H_cm": round(float(h), 2) if h else None,
        "radio_hemisferio_px": round(float(cal["virtual_radius_px"]), 1) if cal.get("virtual_radius_px") else None,
        "calibracion_foto_px": "{}×{}".format(w, hh) if w and hh else None,
        "calibracion_guardada": cal.get("saved_at"),
        "focal_equivalente_mm": round(fmm, 2) if fmm else None,
    }


CUT_RULE_LABELS = {
    "glama": "recomendado (log-normal) por foto, iterado desde 128; Otsu provisional si no hay recomendación",
    "otsu": "automático: Otsu (1979) por foto",
    "fixed": "fijo (mismo valor para todas las fotos)",
}
DECISIONS_FILE = "CobDosel_revision_cortes.json"


def _decision_summary(results):
    """Texto «20 fotos: 16 automáticas, 2 aceptadas, 1 ajustada, 1 excluida»."""
    from collections import Counter
    cnt = Counter()
    for r in results:
        d = (r.get("decision") or "").split(" (")[0]
        if d:
            cnt[d] += 1
    if not cnt:
        return ""
    order = ("automático", "aceptado", "ajustado", "excluido", "pendiente de revisión")
    labels = {"automático": ("automática", "automáticas"),
              "aceptado": ("revisada sin cambio", "revisadas sin cambio"),
              "ajustado": ("ajustada", "ajustadas"),
              "excluido": ("excluida por el operador", "excluidas por el operador"),
              "pendiente de revisión": ("pendiente de revisión", "pendientes de revisión")}
    return "{} fotos: {}".format(sum(cnt.values()), ", ".join(
        "{} {}".format(cnt[k], labels[k][0 if cnt[k] == 1 else 1]) for k in order if cnt.get(k)))


def load_decisions(folder):
    """Decisiones de la revisión visual guardadas junto a las fotos ({} si no hay archivo)."""
    p = os.path.join(folder, DECISIONS_FILE)
    if not os.path.isfile(p):
        return {}
    with open(p, "r", encoding="utf-8") as f:
        data = json.load(f)
    return dict(data.get("fotos") or {})


def save_decisions(folder, decisions, version=""):
    """Guarda las decisiones (archivo JSON en la carpeta de fotos). Devuelve la ruta."""
    p = os.path.join(folder, DECISIONS_FILE)
    with open(p, "w", encoding="utf-8") as f:
        json.dump({"plugin": "CobDosel", "version": version,
                   "guardado": datetime.datetime.now().isoformat(timespec="seconds"),
                   "fotos": decisions}, f, indent=2, ensure_ascii=False)
    return p


def aggregate(results, planned_plots):
    """Medias por parcela (fotos válidas disponibles) y media del sitio (Índice CaCo)."""
    plots = {}
    order = []
    for r in results:
        p = r["parcela"]
        if p not in plots:
            plots[p] = {"parcela": p, "reemplaza_a": r.get("reemplaza_a") or "", "fotos": 0,
                        "validas": 0, "caco": []}
            order.append(p)
        plots[p]["fotos"] += 1
        if r.get("valida"):
            plots[p]["validas"] += 1
            plots[p]["caco"].append(r["caco"])
        if r.get("reemplaza_a"):
            plots[p]["reemplaza_a"] = r["reemplaza_a"]
    rows, warnings = [], []
    for p in order:
        d = plots[p]
        vals = d.pop("caco")
        d["caco_media"] = float(np.mean(vals)) if vals else None
        d["caco_min"] = float(np.min(vals)) if vals else None
        d["caco_max"] = float(np.max(vals)) if vals else None
        if not vals:
            warnings.append("Parcela {}: sin fotos válidas; no entra en la media del sitio "
                            "(aplique el mecanismo de reemplazo).".format(p))
        rows.append(d)
    means = [d["caco_media"] for d in rows if d["caco_media"] is not None]
    site_mean = float(np.mean(means)) if means else None
    n_ok = len(means)
    if planned_plots and n_ok != int(planned_plots):
        warnings.append("Parcelas con resultado: {} de {} planificadas.".format(n_ok, int(planned_plots)))
    site = {
        "parcelas_planificadas": int(planned_plots) if planned_plots else None,
        "parcelas_con_resultado": n_ok,
        "fotos_validas": sum(d["validas"] for d in rows),
        "fotos_total": sum(d["fotos"] for d in rows),
        "caco_sitio": site_mean,
        "completo": bool(planned_plots) and n_ok == int(planned_plots),
        "pendientes_revision": sum(1 for r in results if r.get("valida") and r.get("revisar")),
        "revision": _decision_summary(results),
    }
    if site["pendientes_revision"]:
        warnings.insert(0, "Resultado provisional: {} foto(s) pendientes de revisión del corte "
                           "(«Revisar cortes…»).".format(site["pendientes_revision"]))
    return {"parcelas": rows, "sitio": site, "avisos": warnings}


# ----------------------------------------------------------------------
# Salidas
# ----------------------------------------------------------------------
PHOTO_CSV_COLUMNS = ("parcela", "reemplaza_a", "punto", "archivo", "id_foto", "emparejado_por",
                     "hora_formulario", "hora_exif", "dif_hora_s", "calidad", "valida", "motivo",
                     "camara", "tamano_px",
                     "mascara_horizonte", "cenit_max", "corte", "metodo_corte", "corte_auto", "decision",
                     "operador", "nota", "motivo_revision", "inicios",
                     "caco", "modif_caco", "apertura", "cierre", "fraccion_claros", "visible_pct")
PLOT_CSV_COLUMNS = ("parcela", "reemplaza_a", "fotos", "validas", "caco_media", "caco_min", "caco_max")


def write_csvs(batch, out_prefix, version=""):
    """Escribe <prefijo>_fotos.csv, <prefijo>_parcelas.csv y <prefijo>_sitio.csv. Devuelve las rutas."""
    paths = []
    p1 = out_prefix + "_fotos.csv"
    with open(p1, "w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f)
        w.writerow(PHOTO_CSV_COLUMNS)
        for r in batch["fotos"]:
            w.writerow([r.get(k) for k in PHOTO_CSV_COLUMNS])
    paths.append(p1)
    p2 = out_prefix + "_parcelas.csv"
    with open(p2, "w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f)
        w.writerow(PLOT_CSV_COLUMNS)
        for r in batch["parcelas"]:
            w.writerow([r.get(k) for k in PLOT_CSV_COLUMNS])
    paths.append(p2)
    p3 = out_prefix + "_sitio.csv"
    with open(p3, "w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f)
        s = batch["sitio"]
        prm = batch["parametros"]
        keys = list(s.keys()) + ["param_" + k for k in prm.keys()] + ["plugin", "plugin_version"]
        w.writerow(keys)
        w.writerow(list(s.values()) + list(prm.values()) + ["CobDosel", version])
    paths.append(p3)
    return paths


def write_html(batch, path, version="", site_name=""):
    """Informe HTML autocontenido del lote (sitio, parcelas, fotos, avisos y parámetros)."""
    import html as _h
    from ._version import fmt_num

    def _f(v, d=2):
        return "—" if v is None else fmt_num(v, d)

    s = batch["sitio"]
    prm = batch["parametros"]
    plot_rows = "".join(
        "<tr><td>{}</td><td>{}</td><td class='n'>{} / {}</td><td class='n'><b>{}</b></td>"
        "<td class='n'>{}</td><td class='n'>{}</td></tr>".format(
            _h.escape(str(p["parcela"])),
            _h.escape("reemplazo de " + str(p["reemplaza_a"])) if p.get("reemplaza_a") else "",
            p["validas"], p["fotos"], _f(p["caco_media"]), _f(p["caco_min"]), _f(p["caco_max"]))
        for p in batch["parcelas"])
    photo_rows = "".join(
        "<tr{}><td>{}</td><td>{}</td><td>{}</td><td>{}</td><td class='n'>{}</td><td>{}</td><td>{}</td><td class='n'>{}</td>"
        "<td class='n'>{}</td><td class='n'>{}</td><td>{}</td></tr>".format(
            "" if r.get("valida") else " class='bad'",
            _h.escape(str(r["parcela"])), _h.escape(str(r["punto"])), _h.escape(str(r["archivo"])),
            _h.escape(str(r["emparejado_por"])), r.get("corte", "—"),
            _h.escape(str(r.get("metodo_corte") or "—")),
            _h.escape(str(r.get("decision") or "")) + (
                " — " + _h.escape(str(r["operador"])) if r.get("operador") else ""),
            _f(r.get("mascara_horizonte"), 1),
            _f(r.get("visible_pct"), 1), _f(r.get("caco")), _h.escape(str(r.get("motivo") or "")))
        for r in batch["fotos"])
    pending = s.get("pendientes_revision") or 0
    pending_html = ("<p style='color:#9a3412;font-weight:700'>Resultado provisional: {} foto(s) pendientes de "
                    "revisión del nivel de corte.</p>".format(pending)) if pending else ""
    warn = "".join("<li>{}</li>".format(_h.escape(w)) for w in batch["avisos"]) or "<li>Ninguno.</li>"
    params = "".join("<tr><td>{}</td><td>{}</td></tr>".format(_h.escape(k), _h.escape(str(v)))
                     for k, v in prm.items())
    doc = """<!DOCTYPE html><html lang="es"><head><meta charset="utf-8"/>
<title>CobDosel — lote {site}</title>
<style>
body{{font-family:Segoe UI,Roboto,Arial,sans-serif;background:#f4f7f6;color:#1f2933;margin:0;padding:24px}}
.c{{max-width:1080px;margin:0 auto}} header{{background:#0b4f5e;color:#fff;padding:18px 24px;border-radius:10px}}
section{{background:#fff;border:1px solid #d9e2ec;border-radius:10px;padding:16px 20px;margin-top:16px}}
.big{{font-size:2.4rem;font-weight:800}}
table{{width:100%;border-collapse:collapse;font-size:.9rem}}
th,td{{padding:6px 8px;border-bottom:1px solid #e4e7eb;text-align:left}}
th{{background:#f0f4f8}} td.n{{text-align:right;font-variant-numeric:tabular-nums}} tr.bad td{{color:#9a3412}}
</style></head><body><div class="c">
<header><h1>CobDosel v{version} — Cobertura del dosel por sitio</h1>
<div>Sitio: {site} · Basado en Tichý (2016), doi:10.1111/jvs.12350</div></header>
<section><div>Cobertura del dosel del sitio (media de las parcelas, Índice CaCo)</div>
<div class="big">{mean} %</div>
<p>Parcelas con resultado: {nok} de {nplan} planificadas · Fotos válidas: {fv} de {ft}</p>
<p>Método del nivel de corte: {rule}</p>
<p>Revisión del nivel de corte: {rev}</p>{pending_html}</section>
<section><h2>Parcelas</h2><table><tr><th>Parcela</th><th>Reemplazo</th><th class='n'>Fotos válidas</th>
<th class='n'>CaCo media (%)</th><th class='n'>Mín.</th><th class='n'>Máx.</th></tr>{plots}</table></section>
<section><h2>Fotos</h2><table><tr><th>Parcela</th><th>Punto</th><th>Archivo</th><th>Emparejado por</th>
<th class='n'>Corte</th><th>Método del corte</th><th>Decisión</th><th class='n'>Máscara (°)</th><th class='n'>Visible (%)</th><th class='n'>CaCo (%)</th>
<th>Observaciones</th></tr>{photos}</table></section>
<section><h2>Avisos</h2><ul>{warn}</ul></section>
<section><h2>Parámetros</h2><table>{params}</table></section>
<footer style="text-align:center;color:#52606d;margin-top:16px">Generado por CobDosel v{version} —
Jorge Fallas (jfallas56@gmail.com)</footer></div></body></html>""".format(
        site=_h.escape(site_name or "—"), version=version, mean=_f(s["caco_sitio"]),
        rev=_h.escape(s.get("revision") or "—"), pending_html=pending_html,
        rule=_h.escape(CUT_RULE_LABELS.get(prm.get("regla_corte"), str(prm.get("regla_corte")))),
        nok=s["parcelas_con_resultado"], nplan=s["parcelas_planificadas"] or "—",
        fv=s["fotos_validas"], ft=s["fotos_total"], plots=plot_rows, photos=photo_rows, warn=warn,
        params=params)
    with open(path, "w", encoding="utf-8") as f:
        f.write(doc)
    return path
